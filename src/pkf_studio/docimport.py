"""Requirement documents → text, in a separate process with limits (python -m pkf_studio.docimport <file> <out.json>).

.txt / .md  as UTF-8 (other encodings replaced)
.docx       paragraphs and tables (python-docx); hidden text, comments and tracked changes are flagged
.pdf        text per page (pypdf); a PDF without text (scanned) is refused
Limits: 10 MB file, 200 pages, 50 MB uncompressed .docx, 60 s.
"""
from __future__ import annotations

import hashlib
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

MAX_BYTES = 10 * 1024 * 1024
MAX_PAGES = 200
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024
TYPES = (".txt", ".md", ".docx", ".pdf")


class ImportFailed(Exception):
    pass


def safe_name(name: str) -> tuple[str, str]:
    """('original.<ext>', ext) — the stored name never comes from the uploaded name."""
    ext = Path(name or "").suffix.lower()
    if ext not in TYPES:
        raise ImportFailed(f"unsupported file type '{ext or '?'}' (use {', '.join(TYPES)})")
    return f"original{ext}", ext


def _docx(path: Path) -> tuple[str, list[str]]:
    flags = []
    with zipfile.ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist()) > MAX_DOCX_UNCOMPRESSED:
            raise ImportFailed("the .docx expands to more than 50 MB")
        names = set(z.namelist())
        xml = z.read("word/document.xml").decode("utf-8", errors="replace") if "word/document.xml" in names else ""
    if "word/comments.xml" in names:
        flags.append("the document has comments (not imported; check them)")
    if re.search(r"<w:(ins|del)\b", xml):
        flags.append("the document has tracked changes (accept or reject them first)")
    if re.search(r"<w:vanish\s*/>", xml):
        flags.append("the document has hidden text (it is imported; check it)")
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts), flags


def _pdf(path: Path) -> tuple[str, list[str]]:
    from pypdf import PdfReader

    r = PdfReader(str(path))
    if len(r.pages) > MAX_PAGES:
        raise ImportFailed(f"the PDF has more than {MAX_PAGES} pages")
    text = "\n".join((p.extract_text() or "") for p in r.pages)
    if len(text.strip()) < 20:
        raise ImportFailed("the PDF has no text layer (scanned?); export it with text or paste the text")
    return text, []


def extract(path: Path) -> dict:
    path = Path(path)
    if path.stat().st_size > MAX_BYTES:
        raise ImportFailed("the file is larger than 10 MB")
    ext = path.suffix.lower()
    if ext in (".txt", ".md"):
        text, flags = path.read_bytes().decode("utf-8", errors="replace"), []
    elif ext == ".docx":
        text, flags = _docx(path)
    elif ext == ".pdf":
        text, flags = _pdf(path)
    else:
        raise ImportFailed(f"unsupported file type '{ext}'")
    text = text.replace("\r\n", "\n").strip()
    if not text:
        raise ImportFailed("the document has no text")
    return {"text": text, "flags": flags, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "chars": len(text)}


def extract_isolated(path: Path, timeout: int = 60) -> dict:
    """Run extract() in a child process (parsers never run inside the app)."""
    out = Path(str(path) + ".extract.json")
    try:
        p = subprocess.run([sys.executable, "-I", "-X", "utf8", "-m", "pkf_studio.docimport", str(path), str(out)],
                           capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout)
    except subprocess.TimeoutExpired:
        raise ImportFailed(f"reading the document took longer than {timeout}s") from None
    if not out.exists():
        raise ImportFailed((p.stderr or "the document could not be read").strip()[-500:])
    doc = json.loads(out.read_text(encoding="utf-8"))
    out.unlink(missing_ok=True)
    if doc.get("error"):
        raise ImportFailed(doc["error"])
    return doc


if __name__ == "__main__":
    src, dst = Path(sys.argv[1]), Path(sys.argv[2])
    try:
        res = extract(src)
    except ImportFailed as exc:
        res = {"error": str(exc)}
    except Exception as exc:  # noqa: BLE001
        res = {"error": f"the document could not be read: {type(exc).__name__}: {exc}"[:500]}
    dst.write_text(json.dumps(res), encoding="utf-8", newline="\n")
