"""pkf Studio pure parts: gate evaluation, roster feasibility, scans, document import, jobs and stale locks, ledger
hashes, branch-rule checks, diagram escaping, capability summary, Terraform rendering and guards."""
from __future__ import annotations

import io
import json
import os
import shutil
import sys
import textwrap
import time
import zipfile
from pathlib import Path

import pytest

from pkf_studio import capabilities, checks, diagram, docimport, infra, jobs, scan
from pkf_studio import ledger as L
from pkf_studio.gates import Evidence, PRView, evaluate, parse_comment, roles_named
from pkf_studio.project import Person, Team, feasibility

ROOT = Path(__file__).resolve().parents[1]
ALL = ["product_owner", "business_analyst", "architect", "data_engineer"]


def team(*people):
    return Team(people=[Person(name=n, email=f"{n}@x.com", roles=r, github=n, github_id=i) for i, (n, r) in enumerate(people, 1)])


def rv(uid, login, state="APPROVED", sha="a" * 40, body="", i=None, t="2026-01-01T00:00:00Z"):
    return Evidence("review", i or uid * 10, uid, login, state, sha, body, t)


def cm(uid, login, body, i=None, edited=False, can_write=True):
    kind, arg = parse_comment(body)
    return Evidence("comment", i or uid * 100, uid, login, kind, arg if kind == "APPROVE" else None, body, "2026-01-02T00:00:00Z", edited,
                    None, can_write)


PR = PRView(number=1, head_sha="a" * 40, author_id=9, author_login="author")


# ---------------------------------------------------------------------------------------------- gates
def test_dev_one_person_approves_every_role_by_comment():
    t = team(("solo", ALL))
    pr = PRView(number=1, head_sha="a" * 40, author_id=1, author_login="solo", commit_user_ids={1})
    r = evaluate("G1", "dev", t, pr, [cm(1, "solo", "/approve aaaaaaa")])
    assert r.passed and {a.role for a in r.approvals} == {"product_owner", "business_analyst", "architect"}
    assert not evaluate("G1", "dev", t, pr, [cm(1, "solo", "/approve bbbbbbb")]).passed       # not the head
    assert not evaluate("G1", "dev", t, pr, [cm(1, "solo", "/approve aaaaaaa", edited=True)]).passed
    assert not evaluate("G1", "dev", t, pr, [cm(1, "solo", "/approve aaaaaaa", can_write=False)]).passed
    r = evaluate("G1", "dev", t, pr, [cm(1, "solo", "/approve aaaaaaa", i=1), cm(1, "solo", "/reject the bucket is wrong", i=2)])
    assert not r.passed and "the bucket is wrong" in r.summary()


def test_strict_needs_distinct_non_author_people_and_reviews_only():
    t = team(("po", ["product_owner"]), ("ba", ["business_analyst"]), ("arch", ["architect", "product_owner"]), ("de", ["data_engineer"]))
    ev = [rv(1, "po"), rv(2, "ba"), rv(3, "arch")]
    assert evaluate("G1", "sit", t, PR, ev).passed
    assert not evaluate("G1", "sit", t, PR, [cm(1, "po", "/approve aaaaaaa"), rv(2, "ba"), rv(3, "arch")]).passed  # comments don't count
    r = evaluate("G1", "sit", t, PR, [rv(2, "ba"), rv(3, "arch")])
    assert not r.passed and r.missing  # arch can fill PO or Architect, not both
    pr = PRView(number=1, head_sha="a" * 40, author_id=9, author_login="x", commit_user_ids={3})   # arch committed
    assert not evaluate("G1", "sit", t, pr, ev).passed
    assert not evaluate("G1", "sit", t, PRView(1, "a" * 40, 9, "x", unlinked_commit=True), ev).passed
    old = [rv(1, "po", sha="b" * 40), rv(2, "ba"), rv(3, "arch")]
    assert not evaluate("G1", "sit", t, PR, old).passed
    assert evaluate("G1", "sit", t, PR, old, equivalent=lambda s: s == "b" * 40).passed       # only base merges since
    later = ev + [rv(1, "po", "CHANGES_REQUESTED", body="split is wrong", i=99, t="2026-01-03T00:00:00Z")]
    r = evaluate("G1", "sit", t, PR, later)
    assert not r.passed and r.blocked and "split is wrong" in r.blocked[0]
    assert not evaluate("G1", "sit", t, PRView(1, "a" * 40, 9, "x", touched_protected=[".github/workflows/pkf-gate.yml"]), ev).passed


def test_approve_as_restricts_roles():
    t = team(("ann", ["product_owner", "architect"]), ("ba", ["business_analyst"]), ("arc", ["architect"]))
    assert roles_named("LGTM\nApprove as: po") == {"product_owner"}
    ev = [rv(1, "ann", body="Approve as: architect"), rv(2, "ba"), rv(3, "arc")]
    r = evaluate("G1", "sit", t, PR, ev)
    assert not r.passed and "Product Owner" in r.missing      # ann limited herself to architect
    assert evaluate("G1", "sit", t, PR, [rv(1, "ann", body="Approve as: po"), rv(2, "ba"), rv(3, "arc")]).passed


def test_feasibility_names_missing_roles():
    assert not feasibility(team(("solo", ALL)), "dev")
    probs = feasibility(team(("solo", ALL)), "sit")
    assert any(p.startswith("G1") for p in probs)
    assert not feasibility(team(("po", ["product_owner"]), ("ba", ["business_analyst"]), ("ar", ["architect"]), ("de", ["data_engineer"])), "prd")


# ---------------------------------------------------------------------------------------------- scans and documents
def test_scan_blocks_secrets_and_flags_personal_data():
    text = "bucket: x\npassword: hunter22\nAKIAABCDEFGHIJKLMNOP\ncontact jane@corp.com or +44 20 7946 0958\nUSUBJID: 1001\n"
    f = scan.scan(text, {"dev@example.com"})
    assert {x.what for x in scan.secrets_only(f)} >= {"password", "AWS access key id"}
    assert {x.what for x in f if x.kind == "pii"} >= {"e-mail address", "phone number", "subject/patient identifier"}
    assert not [x for x in scan.scan("mail dev@example.com", {"dev@example.com"})]
    assert all("hunter22" not in x.excerpt for x in f)


def test_docimport_txt_docx_pdf(tmp_path):
    t = tmp_path / "r.txt"
    t.write_text("TRIGGER\nEvery day\n", encoding="utf-8")
    assert docimport.extract_isolated(t)["text"].startswith("TRIGGER")
    import docx

    d = docx.Document()
    d.add_paragraph("SOURCE: the bucket")
    tb = d.add_table(rows=1, cols=2)
    tb.cell(0, 0).text, tb.cell(0, 1).text = "rule", "CDAP"
    run = d.add_paragraph().add_run("secret plan")
    run.font.hidden = True
    p = tmp_path / "r.docx"
    d.save(p)
    out = docimport.extract_isolated(p)
    assert "SOURCE: the bucket" in out["text"] and "rule | CDAP" in out["text"]
    assert any("hidden text" in f for f in out["flags"])
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(200, 200)
    pdf = tmp_path / "scan.pdf"
    with open(pdf, "wb") as fh:
        w.write(fh)
    with pytest.raises(docimport.ImportFailed, match="no text layer"):
        docimport.extract_isolated(pdf)
    bomb = tmp_path / "bomb.docx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "0" * (51 * 1024 * 1024))
    with pytest.raises(docimport.ImportFailed, match="50 MB"):
        docimport.extract_isolated(bomb)
    with pytest.raises(docimport.ImportFailed):
        docimport.safe_name("evil.exe")
    assert docimport.safe_name("../../x.DOCX") == ("original.docx", ".docx")


# ---------------------------------------------------------------------------------------------- jobs
def test_jobs_lock_stale_and_inline(tmp_path, monkeypatch):
    from pkf_studio import actions

    calls = []

    def hello(ctx, who="x"):
        ctx.log(f"hello {who} password: hunter22")
        calls.append(who)
        return {"said": who}

    monkeypatch.setitem(actions.ACTIONS, "hello", hello)
    monkeypatch.setenv("PKF_STUDIO_INLINE_JOBS", "1")
    jid = jobs.start(tmp_path, "p1", "hello", {"who": "ann"})
    st = jobs.status(tmp_path, jid)
    assert st["state"] == "succeeded" and st["result"]["said"] == "ann" and calls == ["ann"]
    assert "hunter22" not in jobs.log(tmp_path, jid) and "[REDACTED]" in jobs.log(tmp_path, jid)
    assert jobs.current(tmp_path, "p1") is None
    # a lock whose runner died is released
    d = jobs.jobs_dir(tmp_path) / "dead"
    d.mkdir(parents=True)
    (d / "cmd.json").write_text(json.dumps({"action": "hello", "pipeline": "p2", "started_at": "2020-01-01T00:00:00+00:00", "root": str(tmp_path)}))
    (d / "pid").write_text("999999 1.0")
    lk = tmp_path / ".pkf-studio" / "locks" / "p2.lock"
    lk.parent.mkdir(parents=True, exist_ok=True)
    lk.write_text(json.dumps({"job": "dead"}))
    assert jobs.current(tmp_path, "p2") is None and not lk.exists()
    assert jobs.status(tmp_path, "dead")["state"] == "crashed"


def test_detached_job_survives_and_reports(tmp_path, monkeypatch):
    monkeypatch.delenv("PKF_STUDIO_INLINE_JOBS", raising=False)
    jid = jobs.start(tmp_path, "p3", "nope")         # unknown action: the runner reports it
    for _ in range(100):
        if jobs.status(tmp_path, jid)["state"] not in ("running", "starting"):
            break
        time.sleep(0.1)
    st = jobs.status(tmp_path, jid)
    assert st["state"] == "failed" and "KeyError" in st["result"]["error"]
    assert jobs.current(tmp_path, "p3") is None


# ---------------------------------------------------------------------------------------------- ledger
def test_input_hash_tracks_what_was_approved():
    files = {"pipelines/p.yaml": "requirement: move files\nexamples: []\nconnection: c\ntrigger: t\n",
             "connections/c.yaml": "source: s\ndestination: d\n", "sources/s.yaml": "requirement: list\n",
             "sources/s.lock.yaml": "plan_digest: abc\n", "triggers/t.yaml": "schedule: '@daily'\n", "destinations/d.yaml": "x: 1\n"}
    read = files.get
    h1 = {g: L.input_hash(g, read, "p") for g in ("G1", "G3", "G4")}
    files["pipelines/p.yaml"] += "owner: other\n"           # not part of G1
    assert L.input_hash("G1", read, "p") == h1["G1"]
    files["pipelines/p.yaml"] = files["pipelines/p.yaml"].replace("move files", "copy files")
    assert L.input_hash("G1", read, "p") != h1["G1"]
    files["sources/s.lock.yaml"] = "plan_digest: def\n"
    assert L.input_hash("G3", read, "p") != h1["G3"]
    files["sources/s.py"] = "def f(): pass\n"
    assert L.input_hash("G4", read, "p") != L.input_hash("G3", read, "p")
    assert L.components(read, "p") == [("trigger", "t"), ("source", "s"), ("destination", "d")]


# ---------------------------------------------------------------------------------------------- checks, diagram, capabilities
def test_branch_rules_check():
    good = [{"type": "pull_request", "parameters": {"required_approving_review_count": 1, "dismiss_stale_reviews_on_push": True,
                                                     "require_code_owner_review": True, "require_last_push_approval": True}},
            {"type": "required_status_checks", "parameters": {"required_status_checks": [{"context": "pkf-gate"}]}},
            {"type": "non_fast_forward"}, {"type": "deletion"}]
    assert checks.branch_rules_ok(good) == []
    missing = checks.branch_rules_ok(good[:1])
    assert "required status check `pkf-gate`" in missing and "block force pushes" in missing


def test_diagram_escapes_labels():
    d = diagram.dot('p"x', [{"kind": "trigger", "name": "t", "steps": [{"id": "a", "op": "sqs", "title": 'say "hi" <b>{x}'}],
                             "schedule": "*/5 * * * *"},
                            {"kind": "source", "name": "s", "steps": []}, {"kind": "destination", "name": "d", "steps": []}])
    assert 'say \\"hi\\" ‹b›(x)' in d and "trigger -> source" in d and "source -> destination" in d
    assert d.count("{") == d.count("}")


def test_capability_summary_flags_deletes_and_network():
    s = capabilities.summarize("import boto3, requests\nimport subprocess\ndef f(aws):\n    aws.client('s3').delete_object(Bucket='b', Key='k')\n"
                               "    open('/etc/passwd')\n")
    assert s["deletes"] == ["delete_object"] and s["network"] == ["requests"] and "subprocess" in s["processes"] and s["file_opens"] == 1


# ---------------------------------------------------------------------------------------------- infra
@pytest.fixture()
def specs(tmp_path):
    dst = tmp_path / "specs"
    shutil.copytree(ROOT / "examples", dst, ignore=shutil.ignore_patterns("dags", "__pycache__", ".git", ".llm"))
    t = dst / "triggers" / "partner_drop_sqs.yaml"
    t.write_text(t.read_text().replace("  queue_url: https://sqs.eu-west-1.amazonaws.com/123456789012/partner-drop-events",
                                       "  queue_url: https://sqs.eu-west-1.amazonaws.com/123456789012/pkf-dev-partner-zip-copy"))  # value only
    os.chdir(dst)
    return dst


def test_terraform_rendering_and_policy(specs, monkeypatch):
    import hcl2

    monkeypatch.setenv("PKF_PLUGIN_ROOT", str(ROOT / "steps-plugin"))
    f = infra.facts(specs, "partner_zip_copy", "dev", "eu-west-1", "123456789012")
    assert f["queue"]["managed"] and f["queue"]["max_receive_count"] == 288      # 24 h at */5
    assert f["notifications"] == [{"bucket": "partner-drop-bucket", "queue_arn": "arn:aws:sqs:eu-west-1:123456789012:pkf-dev-partner-zip-copy",
                                   "prefix": "incoming/", "suffix": ".zip", "pipeline": "partner_zip_copy"}]
    man = infra.render(specs, "partner_zip_copy", f, "dev1", "acme/specs", {"cost-centre": "rd"}, {"bucket": "tf-state", "region": "eu-west-1"})
    assert man["needs_apply"] and man["buckets"] == ["partner-drop-bucket"]
    for tf in [*(specs / "infra").rglob("*.tf")]:
        hcl2.loads(tf.read_text())                                                # well-formed HCL
    tfv = json.loads((specs / "infra/partner_zip_copy/envs/dev/terraform.tfvars.json").read_text())
    assert tfv["name"] == "pkf-dev-partner-zip-copy" and tfv["tags"]["cost-centre"] == "rd" and tfv["tags"]["managed_by"] == "pkf-studio"
    assert "use_lockfile = true" in (specs / "infra/partner_zip_copy/envs/dev/backend.hcl").read_text()
    assert "<REQUIRED" in (specs / "infra/partner_zip_copy/envs/prd/terraform.tfvars.json").read_text()
    pol = json.loads((specs / "infra/partner_zip_copy/iam-policy.json").read_text())
    acts = {a for s in pol["Statement"] for a in s["Action"]}
    assert {"sqs:ReceiveMessage", "sqs:DeleteMessage", "s3:GetObject", "s3:PutObject", "s3:ListBucket"} <= acts
    bucket_tf = (specs / "infra/buckets/partner-drop-bucket/main.tf").read_text()
    assert '"pkf-partner_zip_copy"' in bucket_tf and 'filter_suffix = ".zip"' in bucket_tf
    # scheduled pipeline: no resources
    f2 = infra.facts(specs, "sas_folder_restructure", "dev", "eu-north-1", None)
    m2 = infra.render(specs, "sas_folder_restructure", f2, "dev1", "acme/specs", {}, {})
    assert not m2["needs_apply"] and m2["note"].startswith("nothing to apply")
    pol2 = json.loads((specs / "infra/sas_folder_restructure/iam-policy.json").read_text())
    assert "s3:DeleteObject" in {a for s in pol2["Statement"] for a in s["Action"]}  # a move deletes the source


def test_plan_guard_refuses_destroy_and_unknown_types():
    pj = {"resource_changes": [{"address": "a", "type": "aws_sqs_queue", "change": {"actions": ["create"]}},
                               {"address": "b", "type": "aws_sqs_queue", "change": {"actions": ["delete", "create"]}},
                               {"address": "c", "type": "aws_iam_role", "change": {"actions": ["create"]}}]}
    probs = infra.plan_guard(pj)
    assert any("replace" in p for p in probs) and any("aws_iam_role" in p for p in probs) and len(probs) == 2


FAKE_TF = textwrap.dedent('''\
    #!{py}
    import json, sys, pathlib
    a = sys.argv[1:]
    if a[0] == "version":
        print(json.dumps({{"terraform_version": "1.10.5"}})); sys.exit(0)
    if a[0] == "init":
        sys.exit(0)
    if a[0] == "plan":
        out = [x.split("=", 1)[1] for x in a if x.startswith("-out=")][0]
        pathlib.Path(out).write_text("PLAN " + str(pathlib.Path.cwd()))
        print("Plan: 4 to add, 0 to change, 0 to destroy."); sys.exit(0)
    if a[0] == "show":
        if "-json" in a:
            print(json.dumps({{"resource_changes": [{{"address": "module.pipeline.aws_sqs_queue.queue[0]", "type": "aws_sqs_queue",
                                                    "change": {{"actions": ["create"]}}}}]}}))
        else:
            print("  # module.pipeline.aws_sqs_queue.queue[0] will be created")
        sys.exit(0)
    if a[0] == "apply":
        print("Apply complete! Resources: 4 added."); sys.exit(0)
    if a[0] == "output":
        print(json.dumps({{"queue_url": {{"value": "https://sqs.eu-west-1.amazonaws.com/123456789012/q"}}}})); sys.exit(0)
''')


def fake_terraform(tmp_path) -> Path:
    p = tmp_path / "terraform"
    p.write_text(FAKE_TF.format(py=sys.executable))
    p.chmod(0o755)
    return p


def test_plan_and_apply_only_the_approved_plan(specs, tmp_path, monkeypatch):
    monkeypatch.setenv("PKF_PLUGIN_ROOT", str(ROOT / "steps-plugin"))
    monkeypatch.setenv("PKF_TERRAFORM", str(fake_terraform(tmp_path)))
    f = infra.facts(specs, "partner_zip_copy", "dev", "eu-west-1", "123456789012")
    infra.render(specs, "partner_zip_copy", f, "dev1", "acme/specs", {}, {})
    stack = specs / "infra/partner_zip_copy/envs/dev"
    log = []
    res = infra.plan_stack(specs, stack, "partner_zip_copy/dev", log.append, local_state=True)
    assert res["changes"] == [("module.pipeline.aws_sqs_queue.queue[0]", ["create"])]
    assert (stack / "plan.sha256").read_text().strip() == res["plan_sha256"] and "will be created" in (stack / "plan.txt").read_text()
    assert any("-backend-config=path=" in x for x in log)
    plan_file = specs / ".pkf-studio/terraform/partner_zip_copy/dev/dev.tfplan"
    good = plan_file.read_text()
    plan_file.write_text("TAMPERED")             # (before the apply: once applied, the approved plan is skipped, D3)
    with pytest.raises(infra.InfraError, match="not the one approved"):
        infra.apply_stack(specs, stack, "partner_zip_copy/dev", log.append)
    plan_file.write_text(good)
    out = infra.apply_stack(specs, stack, "partner_zip_copy/dev", log.append)
    assert out["outputs"]["queue_url"].endswith("/q") and out["applied_sha256"] == res["plan_sha256"]


def test_foreign_notifications_detected():
    class S3:
        def get_bucket_notification_configuration(self, Bucket):
            return {"QueueConfigurations": [{"QueueArn": "arn:mine"}, {"QueueArn": "arn:theirs"}], "LambdaFunctionConfigurations": [{"LambdaFunctionArn": "arn:fn"}]}

    assert infra.foreign_notifications(S3(), "b", {"arn:mine"}) == ["arn:theirs", "arn:fn"]


def test_schedule_minutes_and_receive_count():
    assert infra.schedule_minutes("*/45 * * * *") == 45 and infra.max_receive_count("*/45 * * * *") == 32
    assert infra.max_receive_count("@daily") == 3 and infra.max_receive_count("*/1 * * * *") == 1000


def test_runner_block_keeps_popen_subclassable():
    """asyncio.windows_utils does `class Popen(subprocess.Popen)` at import; the block must allow that."""
    import subprocess as sp, sys, textwrap
    from pkf.design import RUNNER
    probe = RUNNER.split("from pkf_runtime import simulate")[0] + textwrap.dedent("""
        class Sub(subprocess.Popen):
            pass
        for cls in (subprocess.Popen, Sub):
            try:
                cls(["echo"])
                print("NOT BLOCKED")
            except PermissionError:
                print("blocked")
        """)
    out = sp.run([sys.executable, "-c", probe], capture_output=True, text=True, timeout=60)
    assert out.returncode == 0, out.stderr
    assert out.stdout.split() == ["blocked", "blocked"]


# ---------------------------------------------------------------------------------------------- 0.5.0: jobs and runner (K1, K2, E3, B7)
def _ctx(tmp_path, pipeline="p"):
    from pkf_studio.runner import Ctx

    d = tmp_path / "job"
    d.mkdir(exist_ok=True)
    return Ctx(d, {"root": str(tmp_path), "pipeline": pipeline})


def test_problem_text_picks_the_hint_by_code():
    from pkf_studio.runner import problem_text

    doc = {"ok": False, "issues": [{"msg": "case a: KeyError", "code": "test_failed"},
                                   {"msg": "the AI coder stopped at step b: rate limited", "code": "llm_error"}], "data": {}}
    t = problem_text(("code", "destination", "d", "--all"), doc)
    # R-phase1 #9: the AI failure is the one message (the partial code's test failures follow from it)
    assert t == ("pkf code destination d: the AI coder stopped at step b: rate limited — the AI step failed; this is usually "
                 "temporary: run it again (Regenerate (AI) / Run tests & build again)")
    notrun = problem_text(("test", "destination", "d"), {"ok": False, "issues": [
        {"msg": "missing custom step function(s) ['b']", "code": "test_failed"},
        {"msg": "case_x: not run: fix the problems first", "code": "test_failed"}]})
    assert "not run" not in notrun and "missing custom step function(s) ['b']" in notrun
    told = problem_text(("code", "destination", "d"), {"ok": False, "issues": [
        {"msg": "the AI answer did not fit the schema — the AI step failed; this is usually temporary: run it again "
                "(Regenerate (AI) / Run tests & build again)", "code": "llm_schema"}]})
    assert told.count("run it again") == 1, told
    assert "no recorded AI answer" in problem_text(("plan", "source", "s"), {"ok": False, "issues": [{"msg": "x"}], "data": {"code": "no_fixture"}})
    assert "reinstall" in problem_text(("test", "destination", "d"), {"ok": False, "issues": [{"msg": "x", "code": "runner_mismatch"}]})
    warn_only = problem_text(("build", "p"), {"ok": False, "issues": [{"msg": "w1", "severity": "warning"}]})
    assert warn_only == "pkf build p: w1"
    assert "S4 status panel" in problem_text(("build", "p", "--draft"), {"ok": False, "issues": [
        {"msg": "destination 'd' is not ready", "code": "component_not_ready"}, {"msg": "warn", "severity": "warning", "code": "x"}]})


def test_ctx_pkf_records_problems_unless_quiet_or_check(tmp_path, monkeypatch):
    from pkf_studio.runner import ActionError, Ctx

    bad = {"ok": False, "issues": [{"msg": "no code yet", "code": "test_failed"}], "data": {}}
    monkeypatch.setattr(Ctx, "_argv", staticmethod(lambda args: [sys.executable, "-c", f"print({json.dumps(json.dumps(bad))})"]))
    ctx = _ctx(tmp_path)
    assert ctx.pkf("test", "pipeline", "p", cwd=tmp_path, quiet=True)["ok"] is False and ctx.problems == []
    ctx.pkf("test", "pipeline", "p", cwd=tmp_path)
    ctx.pkf("test", "pipeline", "p", cwd=tmp_path)                       # the same problem once
    assert ctx.problems == ["pkf test pipeline p: no code yet — the generated code fails its tests: see the S4 status panel (Workflow)"]
    with pytest.raises(ActionError, match="no code yet"):
        ctx.pkf("test", "pipeline", "p", cwd=tmp_path, check=True)
    assert len(ctx.problems) == 1
    for i in range(40):
        ctx.problem(f"p{i} " + "x" * 700)
    assert len(ctx.problems) == 30 and all(len(p) <= 600 for p in ctx.problems)
    assert "⚠ p0" in (ctx.dir / "log").read_text(encoding="utf-8")


def test_job_result_carries_problems_and_severity(tmp_path, monkeypatch):
    from pkf_studio import actions
    from pkf_studio.runner import ActionError

    def warn(ctx):
        ctx.problem("pkf code destination d: the AI failed")
        return {"done": 1}

    def failed_run(ctx):
        return {"severity": "error", "problems": ["the run failed: 2 files"]}

    def refuse(ctx):
        ctx.problem("pushed? no")
        raise ActionError("not now")

    for name, fn in (("warn", warn), ("failed_run", failed_run), ("refuse", refuse)):
        monkeypatch.setitem(actions.ACTIONS, name, fn)
    monkeypatch.setenv("PKF_STUDIO_INLINE_JOBS", "1")
    st = jobs.status(tmp_path, jobs.start(tmp_path, "p1", "warn"))
    assert st["state"] == "succeeded" and st["result"]["problems"] == ["pkf code destination d: the AI failed"]
    assert st["result"]["severity"] == "warning" and st["result"]["done"] == 1
    st = jobs.status(tmp_path, jobs.start(tmp_path, "p1", "failed_run"))
    assert st["state"] == "succeeded" and st["result"]["severity"] == "error" and st["result"]["problems"] == ["the run failed: 2 files"]
    st = jobs.status(tmp_path, jobs.start(tmp_path, "p1", "refuse"))
    assert st["state"] == "failed" and st["result"]["error"] == "not now" and st["result"]["problems"] == ["pushed? no"]
    monkeypatch.setitem(actions.ACTIONS, "fine", lambda ctx: {"x": 1})
    st = jobs.status(tmp_path, jobs.start(tmp_path, "p1", "fine"))
    assert st["result"]["problems"] == [] and st["result"]["severity"] is None


def test_long_step_logs_heartbeat(tmp_path, monkeypatch):
    from pkf_studio.runner import Ctx

    argv = Ctx._argv
    doc = {"ok": True, "issues": [], "data": {"x": 1}}
    script = f"import time; time.sleep(0.8); print({json.dumps(json.dumps(doc))})"
    monkeypatch.setattr(Ctx, "_argv", staticmethod(lambda args: [sys.executable, "-c", script]))
    monkeypatch.setenv("PKF_STUDIO_HEARTBEAT_S", "0.2")
    ctx = _ctx(tmp_path)
    assert ctx.pkf("code", "destination", "d", cwd=tmp_path) == doc
    log = (ctx.dir / "log").read_text(encoding="utf-8")
    assert "  … still running pkf code destination d (0s)" in log and log.rstrip().endswith("→ OK"), log
    monkeypatch.setattr(Ctx, "_argv", staticmethod(lambda args: [sys.executable, "-c", "import time; time.sleep(30)"]))
    t0 = time.monotonic()
    out = ctx.pkf("code", "destination", "d", cwd=tmp_path, timeout=1)
    assert time.monotonic() - t0 < 15 and not out["ok"] and out["issues"][0]["code"] == "timeout"
    assert ctx.problems == ["pkf code destination d: timed out after 1s"]
    assert argv(("status", "p"))[1:] == ["-P", "-m", "pkf.cli", "--json", "status", "p"]


def test_windows_creationflags_hide_consoles(tmp_path, monkeypatch):
    calls = []

    class FakePopen:
        pid = 4242

        def __init__(self, argv, **kw):
            calls.append(kw.get("creationflags"))
            if len(calls) == 1:
                raise OSError(5, "Access is denied")   # the console's job object forbids breakaway

    monkeypatch.delenv("PKF_STUDIO_INLINE_JOBS", raising=False)
    import types

    fake_os = types.SimpleNamespace(**{k: getattr(os, k) for k in dir(os) if not k.startswith("__")})
    fake_os.name = "nt"                                # only jobs.py sees Windows (pathlib keeps working)
    monkeypatch.setattr(jobs, "os", fake_os)
    monkeypatch.setattr(jobs.subprocess, "Popen", FakePopen)
    jid = jobs.start(tmp_path, "p1", "nope")
    assert calls == [0x08000000 | 0x200 | 0x01000000, 0x08000000 | 0x200]
    assert not calls[0] & 0x8 and not calls[1] & 0x8, "never DETACHED_PROCESS (each console child would get a window)"
    assert "could not leave the console's job object" in jobs.log(tmp_path, jid)
    assert (jobs.jobs_dir(tmp_path) / jid / "pid").read_text().startswith("4242")
    jobs._lock(tmp_path, "p1").unlink()


class _FakeStudio:
    def __init__(self, wt, branch, push_msg=None):
        self.wt, self.branch, self.push_msg, self.pushed = wt, branch, push_msg, []

    def worktree_path(self, p):
        return self.wt

    def worktree_branch(self, p):
        return self.branch

    def push_quiet(self, p):
        self.pushed.append(p)
        return self.push_msg


def test_stage_job_pushes_at_the_end_even_when_it_fails(tmp_path):
    from pkf_studio import actions
    from pkf_studio.runner import ActionError

    wt = tmp_path / "wt"
    (wt / ".git").mkdir(parents=True)

    @actions.stage_job
    def ok(ctx):
        ctx.studio = s
        return {"x": 1}

    @actions.stage_job
    def refused(ctx):
        ctx.studio = s
        raise ActionError("nope")

    @actions.stage_job
    def crashed(ctx):
        ctx.studio = s
        raise KeyError("boom")

    s = _FakeStudio(wt, "pkf/p/g4")
    ctx = _ctx(tmp_path)
    assert ok(ctx) == {"x": 1} and s.pushed == ["p"] and ctx.problems == []
    with pytest.raises(ActionError):
        refused(ctx)
    with pytest.raises(KeyError):
        crashed(ctx)
    assert s.pushed == ["p", "p", "p"]
    s.push_msg = "GitHub's copy of pkf/p/g4 changed (someone pushed or used 'Update branch'): Workflow → GitHub's copy changed"
    ok(ctx)
    assert ctx.problems == [s.push_msg]
    for branch in ("main", "pkf/p/run-20261018t090000", "pkf/other/g4"):   # only the pipeline's own stage branches
        s = _FakeStudio(wt, branch)
        ok(_ctx(tmp_path))
        assert s.pushed == []
    assert ok.__name__ == "ok" and "stage_job" not in actions.ACTIONS


def test_stage_jobs_are_decorated():
    from pkf_studio import actions

    for name in ("analyze", "resplit", "start_g2", "start_g3", "plan_components", "plan_component", "start_g4", "test_build",
                 "code_component", "start_g5", "prepare_env"):
        assert getattr(actions.ACTIONS[name], "__wrapped__", None) is not None, name
    assert not hasattr(actions.ACTIONS["run_local"], "__wrapped__")


def test_capability_summary_names_uncoded_steps(tmp_path):
    text = capabilities.report(tmp_path, [], ["destination d: move_and_tag, compute_target_keys"])
    assert "No generated code yet: destination d has custom steps move_and_tag, compute_target_keys that are not coded " \
           "(Workflow → S4 → Regenerate (AI))" in text
    assert "every component uses built-in operations only" not in text
    assert "built-in operations only" in capabilities.report(tmp_path, [])


def test_run_local_needs_a_preview_for_custom_delivery(monkeypatch, tmp_path):
    """D1 (replaces the Phase-1 checkpoint guard): a pipeline whose destination moves files itself runs only from a
    preview; without one it is refused before any branch is made."""
    from pkf_studio import actions, stages

    files = {"pipelines/p.yaml": "connection: c\ntrigger: t\n", "connections/c.yaml": "source: s\ndestination: d\n",
             "destinations/d.lock.yaml": "plan:\n  steps:\n  - {id: keys, op: custom}\n  - {id: move_and_tag, op: custom_delivery}\n"}
    assert actions._custom_delivery(files.get, "p")
    plain = {**files, "destinations/d.lock.yaml": "plan:\n  steps:\n  - {id: keys, op: custom}\n"}
    assert not actions._custom_delivery(plain.get, "p")

    class S:
        read_base = staticmethod(files.get)

        def prs(self, p, age):
            return {}

        def open_run_branch(self, p, branch=None):
            raise AssertionError("no branch may be made")

    class G:
        gate, done = "G1", True

    monkeypatch.setattr(actions, "_open", lambda ctx: S())
    monkeypatch.setattr(stages, "derive", lambda s, p, prs: {"gates": [G()]})
    with pytest.raises(actions.ActionError, match="this pipeline moves or deletes files: Preview this run first, then Move"):
        actions.run_local(_ctx(tmp_path), own_credentials=True)
    with pytest.raises(actions.ActionError, match="preview's run variables"):
        actions.run_local(_ctx(tmp_path), own_credentials=True, planned=str(tmp_path / "x.json"), date="2026-10-15")
    S.read_base = staticmethod(plain.get)
    with pytest.raises(actions.ActionError, match="use Run locally"):
        actions.run_local(_ctx(tmp_path), own_credentials=True, planned=str(tmp_path / "x.json"))


def test_start_g5_rerecords_a_manual_code_approval_with_g4_evidence(monkeypatch, tmp_path):
    """§6.8: `pkf approve-code --by …` by hand (V1's workaround) has no evidence: S5 records it again with G4's."""
    from pkf_studio import actions

    wt = tmp_path / "wt"
    (wt / "destinations").mkdir(parents=True)
    (wt / "destinations" / "d.lock.yaml").write_text("code: {file: d.py}\ncode_approval: {by: '@me', digest: x}\n")
    (wt / "destinations" / "e.lock.yaml").write_text("code: {file: e.py}\ncode_approval: {by: '@me', digest: x, evidence: {url: u}}\n")
    calls = []

    class S:
        def evidence_args(self, row):
            return ["--by", "@dev1", "--evidence", "https://pr/4"]

    ctx = _ctx(tmp_path)
    monkeypatch.setattr(actions, "_open", lambda c: S())
    monkeypatch.setattr(actions, "_stage", lambda c, s, g: wt)
    monkeypatch.setattr(actions, "_ledger", lambda c, s, g: {"gate": g})
    monkeypatch.setattr(actions, "_commit", lambda c, s, m: None)
    monkeypatch.setattr(actions, "prepare_env_in", lambda c, s, w: {"prepared": True})
    monkeypatch.setattr(actions, "_push_at_end", lambda c: None)
    monkeypatch.setattr(actions, "_status", lambda c, w: {"components": [
        {"kind": "destination", "name": "d", "stage": "approved", "mode": "requirement"},
        {"kind": "destination", "name": "e", "stage": "approved", "mode": "requirement"},
        {"kind": "source", "name": "s", "stage": "code_generated", "mode": "requirement"},
        {"kind": "trigger", "name": "t", "stage": "approved", "mode": "config"}]})
    monkeypatch.setattr(ctx, "pkf", lambda *a, **kw: calls.append(a) or {"ok": True})
    assert actions.start_g5(ctx) == {"prepared": True}
    assert calls == [("approve-code", "destination", "d", "--by", "@dev1", "--evidence", "https://pr/4"),
                     ("approve-code", "source", "s", "--by", "@dev1", "--evidence", "https://pr/4"), ("build", "p")]


# ---------------------------------------------------------------------------------------------- Phase 2 (WP4)
def test_capability_summary_warns_copy_object(tmp_path):
    """B3: a single CopyObject call is limited to 5 GiB; the managed copy (multipart) is not."""
    src = "def move(transfers, params, ctx, aws):\n    aws.s3_destination.copy_object(Bucket='b', Key='k', CopySource={})\n"
    assert capabilities.summarize(src)["warnings"] == [capabilities.COPY_OBJECT_WARNING]
    managed = "def move(transfers, params, ctx, aws):\n    aws.s3_destination.copy({}, 'b', 'k')\n"
    assert capabilities.summarize(managed)["warnings"] == []
    (tmp_path / "d.py").write_text(src)
    assert ("- **Warning:** copy_object is limited to 5 GiB per object; the managed copy (multipart) is needed for larger files"
            in capabilities.report(tmp_path, ["d.py"]))


def test_runner_stray_output_goes_to_crash_log(tmp_path, monkeypatch):
    """B8: whatever the job's process prints outside its log (here a sitecustomize that prints on start-up) lands in
    crash.log; `log` holds only the runner's lines. E12: crash_tail returns its end."""
    monkeypatch.delenv("PKF_STUDIO_INLINE_JOBS", raising=False)
    site = tmp_path / "site"
    site.mkdir()
    (site / "sitecustomize.py").write_text("import sys\nprint('stray stdout line')\nprint('stray stderr line', file=sys.stderr)\n")
    monkeypatch.setenv("PYTHONPATH", str(site))
    root = tmp_path / "root"
    root.mkdir()
    jid = jobs.start(root, "p3", "nope")
    for _ in range(150):
        if jobs.status(root, jid)["state"] not in ("running", "starting"):
            break
        time.sleep(0.1)
    assert jobs.status(root, jid)["state"] == "failed"
    log = (jobs.jobs_dir(root) / jid / "log").read_text(encoding="utf-8")
    assert log.startswith("== ") and "stray" not in log, log
    tail = jobs.crash_tail(root, jid)
    assert "stray stdout line" in tail and "stray stderr line" in tail
    assert jobs.crash_tail(root, jid, n=1) == tail.splitlines()[-1]
    assert jobs.crash_tail(root, "no-such-job") == ""


def test_jobs_start_takes_lock_before_creating_dir(tmp_path, monkeypatch):
    """C19/K11: the lock comes first (two clicks cannot start two jobs), a job folder that cannot be made releases it,
    and the job runs with PKF_STUDIO_JOB=<its id> (restored afterwards for inline jobs)."""
    from pkf_studio import actions

    seen = {}
    real = jobs.jobs_dir

    def spy(root):
        seen.setdefault("lock_first", jobs._lock(root, "p1").exists())
        return real(root)

    def who(ctx):
        seen["env"] = os.environ.get("PKF_STUDIO_JOB")
        return {}

    monkeypatch.setitem(actions.ACTIONS, "who", who)
    monkeypatch.setenv("PKF_STUDIO_INLINE_JOBS", "1")
    monkeypatch.setenv("PKF_STUDIO_JOB", "outer")
    monkeypatch.setattr(jobs, "jobs_dir", spy)
    jid = jobs.start(tmp_path, "p1", "who")
    assert seen == {"lock_first": True, "env": jid} and os.environ["PKF_STUDIO_JOB"] == "outer"
    assert jobs.current(tmp_path, "p1") is None
    # the lock was just taken and the job folder is not there yet: a second start is refused, the lock kept
    lk = jobs._lock(tmp_path, "p1")
    lk.write_text(json.dumps({"job": "20990101T000000-abcdef"}))
    with pytest.raises(jobs.JobError, match="already running"):
        jobs.start(tmp_path, "p1", "who")
    assert jobs.current(tmp_path, "p1") == "20990101T000000-abcdef"
    lk.unlink()
    # the job folder cannot be created: the lock is released again
    monkeypatch.setattr(jobs, "jobs_dir", real)
    (tmp_path / ".pkf-studio" / "jobs").rename(tmp_path / ".pkf-studio" / "jobs-old")
    (tmp_path / ".pkf-studio" / "jobs").write_text("not a folder")
    with pytest.raises(OSError):
        jobs.start(tmp_path, "p1", "who")
    assert not lk.exists()


def test_running_all_lists_every_pipeline(tmp_path):
    """E18: the running jobs of every pipeline (the sidebar)."""
    import psutil

    me = psutil.Process(os.getpid())
    for i, p in enumerate(("alpha", "beta", "gamma")):
        jid = f"20261018T09000{i}-aaaaa{i}"
        d = jobs.jobs_dir(tmp_path) / jid
        d.mkdir(parents=True)
        (d / "cmd.json").write_text(json.dumps({"id": jid, "action": "start_g4", "pipeline": p, "title": "Start S4",
                                                "started_at": "2026-10-18T09:00:00+00:00"}))
        (d / "pid").write_text(f"{os.getpid()} {me.create_time()}")
        if p == "gamma":
            (d / "exit").write_text("0")       # finished: its lock is stale
        lk = jobs._lock(tmp_path, p)
        lk.parent.mkdir(parents=True, exist_ok=True)
        lk.write_text(json.dumps({"job": jid}))
    rows = jobs.running_all(tmp_path)
    assert [(r["pipeline"], r["state"], r["title"]) for r in rows] == [("alpha", "running", "Start S4"), ("beta", "running", "Start S4")]
    assert not jobs._lock(tmp_path, "gamma").exists()
    assert jobs.running_all(tmp_path / "nothing-here") == []


def test_child_env_bedrock_passthrough():
    """Env-2: AWS_* reach the AI commands only when the AI runs on Bedrock; other secrets stay scrubbed."""
    from pkf_studio.runner import child_env

    base = {"AWS_PROFILE": "dev", "AWS_ACCESS_KEY_ID": "AKIA", "AWS_REGION": "eu-north-1", "GITHUB_TOKEN": "gh", "ANTHROPIC_API_KEY": "k",
            "PATH": "/bin"}
    plain = child_env(["code", "destination", "d"], base)
    assert "AWS_PROFILE" not in plain and "AWS_ACCESS_KEY_ID" not in plain and plain["AWS_REGION"] == "eu-north-1"
    assert plain["ANTHROPIC_API_KEY"] == "k" and "GITHUB_TOKEN" not in plain and plain["PYTHONUTF8"] == "1"
    for flag in ("CLAUDE_CODE_USE_BEDROCK", "PKF_LLM_AWS_PASSTHROUGH"):
        env = child_env(["code", "destination", "d"], {**base, flag: "1"})
        assert env["AWS_PROFILE"] == "dev" and env["AWS_ACCESS_KEY_ID"] == "AKIA" and "GITHUB_TOKEN" not in env
        assert "AWS_PROFILE" not in child_env(["test", "destination", "d"], {**base, flag: "1"}), "not for non-AI commands"
        assert "AWS_PROFILE" not in child_env(["build", "p"], {**base, flag: "1"})
    run = child_env(["run", "p"], base)
    assert run["AWS_PROFILE"] == "dev" and run["AWS_ACCESS_KEY_ID"] == "AKIA"
    assert "AWS_PROFILE" not in child_env(["status"], base) and base["AWS_PROFILE"] == "dev", "pure: the input is unchanged"


def test_second_run_skips_the_applied_plan(specs, tmp_path, monkeypatch):
    """D3: once the approved plan is applied, "Run again" skips it (a saved plan cannot be applied twice)."""
    monkeypatch.setenv("PKF_PLUGIN_ROOT", str(ROOT / "steps-plugin"))
    tf = fake_terraform(tmp_path)
    calls = tmp_path / "applies.txt"
    tf.write_text(tf.read_text().replace('if a[0] == "apply":\n',
                                         f'if a[0] == "apply":\n    open({str(calls)!r}, "a").write("apply\\n")\n'))
    monkeypatch.setenv("PKF_TERRAFORM", str(tf))
    f = infra.facts(specs, "partner_zip_copy", "dev", "eu-west-1", "123456789012")
    infra.render(specs, "partner_zip_copy", f, "dev1", "acme/specs", {}, {})
    stack = specs / "infra/partner_zip_copy/envs/dev"
    res = infra.plan_stack(specs, stack, "partner_zip_copy/dev", lambda m: None, local_state=True)
    first = infra.apply_stack(specs, stack, "partner_zip_copy/dev", lambda m: None)
    log = []
    second = infra.apply_stack(specs, stack, "partner_zip_copy/dev", log.append)
    assert calls.read_text().splitlines() == ["apply"], "terraform apply ran once over two runs"
    assert first["applied_sha256"] == res["plan_sha256"] and "outputs" in first
    assert second == {"stack": str(stack), "skipped": "already applied", "applied_sha256": res["plan_sha256"]}
    assert (specs / ".pkf-studio/terraform/partner_zip_copy/dev/applied.sha256").read_text().strip() == res["plan_sha256"]
    assert any("already applied" in m for m in log)
    # a new approved plan (S5 again) is applied
    (stack / "plan.sha256").write_text("0" * 64 + "\n")
    with pytest.raises(infra.InfraError, match="not the one approved"):
        infra.apply_stack(specs, stack, "partner_zip_copy/dev", lambda m: None)


def test_deploy_notes_mention_staging():
    """D6: every Airflow worker must reach the same staging."""
    from pkf_studio import actions

    text = actions.deploy_notes("p", {"needs_apply": False, "note": "nothing to apply — policy only"})
    assert "Set PKF_STAGING_URL=s3://<bucket>/<prefix> on the workers (with a lifecycle rule) unless Airflow runs on one machine" in text


def _prepare_env_world(monkeypatch, tmp_path, dry):
    from pkf_studio import actions

    wt = tmp_path / "wt"
    (wt / "infra" / "p").mkdir(parents=True, exist_ok=True)
    (wt / "infra" / "p" / "iam-policy.json").write_text("{}")
    commits = []

    class Cfg:
        aws_region, aws_account_id, terraform_backend, tags = "eu-north-1", "123456789012", {}, {}

    class S:
        cfg, owner, repo, strict = Cfg(), "acme", "specs", False
        read_base = staticmethod(lambda q: None)

        def me(self):
            return None

    monkeypatch.setattr(infra, "facts", lambda *a: {"warnings": [], "notifications": [], "source": {"type": "local"}})
    monkeypatch.setattr(infra, "repo_notifications", lambda wt: [])
    monkeypatch.setattr(infra, "render", lambda *a: {"warnings": [], "runtime_calls": [], "needs_apply": False, "buckets": [],
                                                     "note": "nothing to apply — policy only"})
    monkeypatch.setattr(infra, "policy_check", lambda *a: {"validate": "not checked", "simulate": "not checked"})
    monkeypatch.setattr(actions, "_commit", lambda c, s, m: commits.append(m))
    ctx = _ctx(tmp_path)
    monkeypatch.setattr(ctx, "pkf", lambda *a, **kw: dry)
    return actions, ctx, S(), wt, commits


def test_prepare_env_records_dry_run_problem(monkeypatch, tmp_path):
    """D9: a failed dry run (e.g. no AWS credentials) is a problem of the S5 job; the S5 files are still committed."""
    dry = {"ok": False, "issues": [{"msg": "NoCredentialsError: Unable to locate credentials", "code": "run_error"}], "data": {}}
    actions, ctx, s, wt, commits = _prepare_env_world(monkeypatch, tmp_path, dry)
    res = actions.prepare_env_in(ctx, s, wt)
    assert res["dry_run_ok"] is False
    assert ctx.problems == ["the dry run failed: NoCredentialsError: Unable to locate credentials"]
    doc = json.loads((wt / "review" / "p" / "dry-run.json").read_text())
    assert doc["ok"] is False and doc["issues"][0]["msg"].startswith("NoCredentialsError")
    assert commits and commits[-1].startswith("S5 p: dev environment rendered") and (wt / "review" / "p" / "deploy.md").exists()
    ok = {"ok": True, "issues": [], "data": {"dry_run": {"files": 2}}}
    actions, ctx, s, wt, commits = _prepare_env_world(monkeypatch, tmp_path, ok)
    assert actions.prepare_env_in(ctx, s, wt)["dry_run_ok"] is True and ctx.problems == []


def test_account_mismatch_refused(monkeypatch, tmp_path):
    """D15/E13: credentials for another account than the pinned one are refused before anything touches AWS."""
    from moto import mock_aws

    from pkf_studio import actions

    for k, v in {"AWS_ACCESS_KEY_ID": "testing", "AWS_SECRET_ACCESS_KEY": "testing", "AWS_DEFAULT_REGION": "eu-north-1"}.items():
        monkeypatch.setenv(k, v)
    monkeypatch.delenv("AWS_PROFILE", raising=False)
    with mock_aws():
        with pytest.raises(infra.InfraError, match="your AWS credentials are for account 123456789012; the project is pinned to 999999999999"):
            infra.check_identity("999999999999", "eu-north-1")
        assert infra.check_identity("123456789012", "eu-north-1")["Account"] == "123456789012"

        files = {"pipelines/p.yaml": "connection: c\ntrigger: t\n", "connections/c.yaml": "source: s\ndestination: d\n",
                 "sources/s.yaml": "type: s3\nconnection: {bucket: b}\n", "destinations/d.yaml": "type: s3\n"}

        class Cfg:
            aws_region, aws_account_id = "eu-north-1", "999999999999"

        class S:
            cfg = Cfg()
            read_base = staticmethod(files.get)

        ctx = _ctx(tmp_path)
        with pytest.raises(actions.ActionError, match="pinned to 999999999999"):
            actions.preflight(ctx, S())
        S.cfg.aws_account_id = "123456789012"
        assert actions.preflight(ctx, S())["Arn"].startswith("arn:aws:")
        assert "AWS credentials: arn:aws:" in (ctx.dir / "log").read_text(encoding="utf-8")
        S.cfg.aws_account_id = None
        ctx = _ctx(tmp_path)
        actions.preflight(ctx, S())
        assert ctx.problems == ["pin the AWS account id in Setup (recommended in dev, required for Terraform)"]
    local = {**files, "sources/s.yaml": "type: local\n", "destinations/d.yaml": "type: local\n"}
    assert not actions.uses_aws(local.get, "p")
    assert actions.uses_aws({**local, "triggers/t.yaml": "params: {queue_url: https://sqs.x/1/q}\n"}.get, "p")


def test_no_credentials_message_has_the_fix(monkeypatch, tmp_path):
    """E13: no (or expired) credentials → the message says how to log in, for the identity check and the run role."""
    for k in ("AWS_ACCESS_KEY_ID", "AWS_SECRET_ACCESS_KEY", "AWS_SESSION_TOKEN", "AWS_PROFILE", "AWS_DEFAULT_PROFILE"):
        monkeypatch.delenv(k, raising=False)
    for k in ("HTTPS_PROXY", "HTTP_PROXY", "https_proxy", "http_proxy", "ALL_PROXY", "all_proxy", "AWS_CONTAINER_CREDENTIALS_FULL_URI", "AWS_CONTAINER_CREDENTIALS_RELATIVE_URI", "AWS_WEB_IDENTITY_TOKEN_FILE"):
        monkeypatch.delenv(k, raising=False)  # the result must not depend on a proxy or container credentials
    monkeypatch.setenv("AWS_SHARED_CREDENTIALS_FILE", str(tmp_path / "none"))
    monkeypatch.setenv("AWS_CONFIG_FILE", str(tmp_path / "none"))
    monkeypatch.setenv("AWS_EC2_METADATA_DISABLED", "true")
    for k in ("AWS_CONTAINER_CREDENTIALS_FULL_URI", "AWS_CONTAINER_CREDENTIALS_RELATIVE_URI", "AWS_WEB_IDENTITY_TOKEN_FILE"):
        monkeypatch.delenv(k, raising=False)
    with pytest.raises(infra.InfraError, match=r"no valid AWS credentials \(NoCredentialsError\): run aws sso login --profile"):
        infra.check_identity("123456789012", "eu-north-1")
    import boto3
    from botocore.exceptions import ClientError

    class STS:
        def assume_role(self, **kw):
            raise ClientError({"Error": {"Code": "ExpiredToken", "Message": "The security token included in the request is expired"}},
                              "AssumeRole")

    monkeypatch.setattr(boto3, "client", lambda *a, **kw: STS())
    with pytest.raises(infra.InfraError, match=r"no valid AWS credentials \(ExpiredToken\): run aws sso login"):
        infra.assume_run_role("arn:aws:iam::123456789012:role/r", {}, "eu-north-1", "p")


def test_files_csv_has_header_and_note(tmp_path):
    """K7: one row per planned file, with the rule or error reason in `note`; LF line ends."""
    from pkf_studio import actions

    rows = [{"package": "p1", "index": 0, "path": "S1/a.csv", "source_key": "incoming/S1/a.csv", "target_key": "S1/a.csv",
             "status": "moved", "detail": "", "note": "rule 7", "size": 3},
            {"package": "p2", "index": 1, "path": "S1/b,c.csv", "source_key": "incoming/S1/b,c.csv", "target_key": None,
             "status": "failed", "detail": "copy failed", "note": "", "size": None}]
    counts = actions.write_files_csv(tmp_path / "files.csv", rows)
    raw = (tmp_path / "files.csv").read_bytes()
    assert b"\r" not in raw and raw.decode().splitlines()[0] == "package,index,path,source_key,target_key,status,detail,note,size"
    import csv

    got = list(csv.DictReader(io.StringIO(raw.decode())))
    assert [r["note"] for r in got] == ["rule 7", ""] and got[1]["path"] == "S1/b,c.csv" and got[1]["target_key"] == ""
    assert counts == {"moved": 1, "failed": 1}


def test_start_warns_about_commits_not_in_the_merged_pr(monkeypatch, tmp_path):
    """C20: commits on a gate's branch made after its PR was merged on GitHub are kept (archive ref) and named."""
    from pkf_studio import actions

    archived = []

    class S:
        def unmerged_local_commits(self, p, gate):
            return 2 if gate == "G2" else 0

        def prs(self, p, age):
            return {"G2": {"number": 7}}

        def archive(self, p, gate):
            archived.append((p, gate))
            return f"refs/pkf-archive/{p}/{gate.lower()}-20261018T120000"

    ctx = _ctx(tmp_path)
    actions._unmerged_warning(ctx, S(), "G1")
    assert ctx.problems == [] and archived == []
    actions._unmerged_warning(ctx, S(), "G2")
    assert archived == [("p", "G2")]
    assert ctx.problems == ["2 commits on g2 were not in the merged PR #7: Reopen G2 to include them "
                            "(they are kept as refs/pkf-archive/p/g2-20261018T120000)"]


def test_failed_start_does_not_publish_an_empty_stage_branch(tmp_path):
    """R-phase1 #5: a Start that fails before its first commit leaves the new stage branch unpublished; a branch with
    work, or one already on GitHub, is pushed as before."""
    from pkf_studio import actions
    from pkf_studio.runner import ActionError

    wt = tmp_path / "wt"
    (wt / ".git").mkdir(parents=True)

    class Git:
        def __init__(self, head, base="b" * 40, remote=None):
            self._head, self._base, self._remote = head, base, remote

        def sha(self, ref):
            return self._base if ref == "origin/main" else self._remote

        def head(self, cwd):
            return self._head

        def is_ancestor(self, a, b, cwd=None):
            return a == self._base

    class S(_FakeStudio):
        base_ref = "origin/main"

    @actions.stage_job
    def start(ctx):
        ctx.studio = s
        raise ActionError("G4 PR #7 was merged without the required approval")

    s = S(wt, "pkf/p/g5")
    s.git = Git(head="b" * 40)                    # the new g5 branch is still the base tip
    ctx = _ctx(tmp_path)
    with pytest.raises(ActionError):
        start(ctx)
    assert s.pushed == [] and "not pushed: pkf/p/g5 has no work of its stage yet" in (ctx.dir / "log").read_text(encoding="utf-8")
    s.git = Git(head="c" * 40)                    # a stage commit beyond the base
    with pytest.raises(ActionError):
        start(ctx)
    assert s.pushed == ["p"]
    s.git = Git(head="b" * 40, remote="b" * 40)   # already on GitHub
    with pytest.raises(ActionError):
        start(ctx)
    assert s.pushed == ["p", "p"]
